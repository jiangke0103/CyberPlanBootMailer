from __future__ import annotations

import argparse
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PLAN_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_ROOT = PLAN_ROOT / "data" / "每日任务"
EXTERNAL_DIR = PLAN_ROOT / "imports"


def load_external_plan(day: date) -> str:
    EXTERNAL_DIR.mkdir(parents=True, exist_ok=True)
    for suffix in (".md", ".txt"):
        path = EXTERNAL_DIR / f"{day.isoformat()}{suffix}"
        if path.exists():
            return path.read_text(encoding="utf-8").strip()
    return ""


COLORS = {
    "primary": "1F4E79", "accent": "2E74B5", "light": "EAF2F8",
    "alt": "F4F8FB", "note": "FFF4CE", "code": "F2F3F5",
    "text": "222222", "muted": "666666",
}

PHASES = [
    (date(2026, 9, 15), date(2026, 9, 20), "启动周", "建立环境、节奏和记录习惯"),
    (date(2026, 9, 21), date(2026, 12, 20), "基础补课与 CET-4", "网络、Linux、Python、安全入门和英语"),
    (date(2026, 12, 21), date(2027, 4, 30), "安全核心与第一个项目", "OWASP、Windows、日志、靶场和作品集"),
    (date(2027, 5, 1), date(2027, 8, 31), "蓝队专项与就业能力", "安全设备、告警研判、应急响应和简历"),
    (date(2027, 9, 1), date(2028, 3, 31), "实习投递与面试", "投递、模拟面试、项目打磨和正式比赛"),
    (date(2028, 4, 1), date(2029, 6, 30), "实习与秋招桥接", "实习、秋招、保底路线和复盘"),
]
WEEKDAYS = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]

NETWORK_TOPICS = [
    "OSI 与 TCP/IP 分层、封装与解封装", "IP 地址、子网划分和网关", "ARP、以太网和交换机基础",
    "VLAN、Trunk 和二层转发", "静态路由和路由表", "动态路由基础（RIP/OSPF 概念）",
    "TCP 三次握手、四次挥手和状态", "UDP、DNS 查询和缓存", "HTTP 请求、响应、状态码和 Cookie",
    "HTTPS、TLS 握手和证书", "DHCP、NAT 和地址转换", "ACL、防火墙规则和最小开放原则",
    "Wireshark 抓包与流量定位", "常见网络排障流程", "VPN 和远程接入基础", "网络知识周测与错题复盘",
]
LINUX_TOPICS = [
    "目录、文件、通配符和帮助命令", "文件权限、用户、组和 umask", "进程、服务、systemd 和日志",
    "软件包管理与软件源", "文本处理 grep、sed、awk 基础", "Shell 变量、条件、循环和脚本",
    "网络命令 ip、ss、curl、ping", "SSH、密钥登录和 scp", "防火墙、端口和服务暴露面",
    "日志定位、时间线和排障", "任务计划、cron 和自动化", "压缩、归档、校验和备份",
    "环境变量、PATH 和权限边界", "Docker 基础与服务容器化", "Linux 安全加固检查", "Linux 综合排障测试",
]
PYTHON_TOPICS = [
    "变量、字符串、数字和输入输出", "列表、元组、字典和集合", "条件、循环和函数", "文件和目录处理",
    "异常处理和日志", "正则表达式和文本提取", "requests 与 HTTP 接口", "JSON、CSV 和数据处理",
    "socket 与网络基础", "subprocess 与系统命令", "Flask 最小 Web 服务", "日志解析小工具",
    "批量巡检小工具", "密码与密钥安全处理", "测试、调试和代码结构", "综合脚本项目复盘",
]
SECURITY_BASE_TOPICS = [
    "CIA、风险、威胁和脆弱性", "身份认证、授权和访问控制", "对称加密、非对称加密和哈希",
    "攻击链、攻击面和最小权限", "OWASP Top 10 总览", "SQL 注入原理与防御", "XSS 原理与防御",
    "文件上传、文件包含和路径穿越", "命令注入、SSRF 和越权", "认证会话和常见逻辑漏洞",
    "Windows 账户、权限和事件日志", "Linux 账户、sudo 和审计日志", "防火墙、IDS、IPS 和 WAF",
    "主机加固、补丁和基线检查", "告警研判和证据链", "应急响应流程与复盘",
]
WEB_TOPICS = [
    "HTTP 请求分析和 Burp 基础", "SQL 注入实验与修复", "XSS 实验与输出编码", "文件上传、文件包含和路径穿越",
    "命令注入和 SSRF", "越权、IDOR 和认证会话", "CSRF、CORS 和安全响应头", "反序列化和组件漏洞基础",
    "信息收集和边界梳理", "漏洞验证和影响判断", "修复建议和复测方法", "Web 安全测试报告结构",
    "OWASP Juice Shop 综合练习", "PortSwigger 综合实验", "Web 漏洞周测", "Web 项目报告复盘",
]
BLUE_TOPICS = [
    "日志源、字段和时间线", "Windows 事件日志分析", "Linux audit 和 syslog 分析", "Web 访问日志攻击痕迹",
    "Wireshark 流量分析", "Zeek/Suricata 告警理解", "Wazuh/ELK 采集与查询", "告警分级和研判",
    "恶意样本基础处置", "主机隔离、取证和恢复", "应急响应演练一", "应急响应演练二",
    "安全设备策略检查", "漏洞修复闭环", "安全运营指标和报表", "蓝队综合项目复盘",
]
CTF_TOPICS = [
    "Linux 命令和文件识别", "编码、进制和隐写基础", "流量包分析入门", "Web 信息收集和源码查看",
    "Web 请求改包和认证绕过", "SQL 注入和 XSS CTF 题", "文件上传、包含和命令执行", "Misc 综合题",
    "Web 综合题", "简单 Crypto 和编码题", "Reverse 基础与字符串分析", "比赛环境和协作流程",
    "写一份完整 Writeup", "复盘一次失败题目", "CTF 周测", "正式比赛或模拟赛",
]


def pick(items: list[str], index: int) -> str:
    return items[index % len(items)]


def phase_for(day: date):
    for start, end, name, focus in PHASES:
        if start <= day <= end:
            return start, end, name, focus
    return PHASES[-1]


def phase_index(day: date) -> int:
    for idx, (start, end, _, _) in enumerate(PHASES):
        if start <= day <= end:
            return idx
    return len(PHASES) - 1


def week_index(day: date, start: date) -> int:
    return max(0, (day - start).days // 7)


def learn_how(output: str) -> str:
    return f"先用 5 分钟回忆上一节，再看 30 分钟资料，随后做 60～80 分钟练习；最后用 10 分钟写下概念、错误和下次动作。卡 10 分钟先记录，卡 30 分钟再查答案。产出：{output}。"


def lab_how(output: str) -> str:
    return f"先写实验目标，再在自有虚拟机或授权靶场复现；记录命令、截图、现象和结论。不要只看教程。做完后故意改一个变量，确认自己理解。产出：{output}。"


def project_how(output: str) -> str:
    return f"只推进一个可验证的小步骤：改配置、运行、验证、提交。每 40 分钟停下来记录一次。产出：{output}。"


def review_how(output: str) -> str:
    return f"按“完成、遗漏、卡点、原因、明天动作”写，不凭感觉打分。产出：{output}。"


def phase0_tasks(day: date):
    data = {
        date(2026, 9, 15): [
            ("09:20-11:20", "搭建 Ubuntu 与 Windows 虚拟机", "安装虚拟机，确认联网、快照和文件复制均正常。", "虚拟机截图和环境说明"),
            ("13:20-15:20", "Linux 常用命令入门", "实操 pwd、ls、cd、mkdir、cp、mv、rm、cat、grep、find。", "20 条命令笔记"),
            ("16:00-17:20", "Python 环境与基础脚本", "运行第一个脚本，练习变量、列表和循环。", "1 个可运行脚本"),
            ("18:40-20:20", "CET-4 诊断和数学竞赛", "计时完成一套听力或阅读 45 分钟，再学数学 1 小时。", "英语错题 5 条、数学题 3 道"),
            ("20:40-21:40", "建立 GitHub 作品集", "创建 network-security-plan 仓库并写 README。", "GitHub 首次提交"),
            ("21:40-22:00", "每日复盘", "记录完成、产出、卡点和明天三件事。", "复盘卡"),
        ],
        date(2026, 9, 16): [
            ("09:20-11:20", "网络基础与子网划分", "学习 OSI/TCP-IP、IP、子网划分并做 10 道题。", "网络笔记和 10 道题"),
            ("13:20-15:20", "Linux 权限与进程", "实操 chmod、chown、ps、top、kill，记录结果。", "权限实验报告"),
            ("16:00-17:20", "HTTP 基础", "理解请求、响应、状态码、Cookie 和 Session。", "一次网页请求的流程图"),
            ("18:40-20:20", "CET-4 与数学竞赛", "英语 45 分钟，数学 1 小时。", "错题本和数学题"),
            ("20:40-21:40", "写解释笔记", "写 800 字：一次网页请求发生了什么。", "技术笔记"),
            ("21:40-22:00", "每日复盘", "按模板复盘。", "复盘卡"),
        ],
        date(2026, 9, 17): [
            ("09:20-11:20", "Wireshark 抓包", "抓取 HTTP 和 DNS，记录源端口、目标端口、请求和响应。", "抓包分析记录"),
            ("13:20-15:20", "Linux 网络排障", "练习 ip、ss、ping、traceroute、curl、tcpdump。", "排障实验记录"),
            ("16:00-17:20", "Python 文本处理", "读取文本、统计行数、筛选关键词。", "1 个日志筛选脚本"),
            ("18:40-20:20", "CET-4 与 CTF 平台", "英语 45 分钟，注册 CTF 平台并查看 Misc 入门题。", "英语错题和平台账号"),
            ("20:40-21:40", "整理抓包结论", "写现象、证据、原因和结论。", "实验报告"),
            ("21:40-22:00", "每日复盘", "按模板复盘。", "复盘卡"),
        ],
        date(2026, 9, 18): [
            ("09:20-11:20", "Nmap 与 Burp 基础", "只扫描自己的虚拟机，学习端口识别和抓包改包。", "工具操作记录"),
            ("13:20-15:20", "安全基础", "理解 CIA、认证、授权和攻击链。", "5 个概念例子"),
            ("16:00-17:20", "安全基础笔记", "用自己的话解释 5 个概念。", "安全笔记"),
            ("18:40-20:20", "CET-4 与复习", "英语 45 分钟，复习本周安全概念。", "错题本"),
            ("20:40-21:40", "输出与提交", "整理本周文件并提交 GitHub。", "GitHub 提交"),
            ("21:40-22:00", "每日复盘", "按模板复盘。", "复盘卡"),
        ],
        date(2026, 9, 19): [
            ("09:20-11:20", "第一次 CTF Misc", "做编码、隐写、文件识别和流量入门题。", "1～2 道题记录"),
            ("13:20-15:20", "CTF Writeup", "写题目、信息、思路、失败尝试和最终解法。", "第一篇 Writeup"),
            ("16:00-17:20", "整理笔记", "把本周笔记按网络、Linux、Python、安全分类。", "笔记目录"),
            ("18:40-20:20", "CET-4 与 GitHub", "英语 45 分钟，提交本周文件。", "GitHub 提交"),
            ("20:40-21:40", "项目骨架", "建立 projects、labs、ctf、notes、reviews 目录。", "作品集结构"),
            ("21:40-22:00", "每日复盘", "按模板复盘。", "复盘卡"),
        ],
        date(2026, 9, 20): [
            ("09:20-11:20", "第一周周测", "网络、Linux、Python、安全各 30 分钟。", "周测分数"),
            ("13:20-15:20", "补最大短板", "只修一个最薄弱点。", "错题闭环"),
            ("16:00-17:20", "完成小项目", "搭建 Linux 静态网页服务器。", "可展示项目"),
            ("18:40-20:00", "CET-4 限时训练", "完成一套题并统计分数。", "分数记录"),
            ("20:20-21:50", "周复盘与下周计划", "按 03_复盘系统.md 完成。", "周复盘"),
            ("21:50-22:00", "收尾", "关闭电脑，准备周一休息。", "周一休息确认"),
        ],
    }
    return data.get(day, [])


def row(time: str, task: str, how: str, output: str):
    return (time, task, how, output)


def make_row(time: str, task: str, mode: str, output: str):
    if mode == "learn":
        how = learn_how(output)
    elif mode == "lab":
        how = lab_how(output)
    elif mode == "project":
        how = project_how(output)
    else:
        how = review_how(output)
    return row(time, task, how, output)


def sample_specs(pidx: int, wd: int, net: str, linux: str, python: str, security: str, web: str, blue: str, ctf: str):
    if pidx == 1:
        specs = {
            1: [(f"计算机网络：{net}", "learn", "1 页概念图 + 10 道练习题"), (f"Linux：{linux}", "lab", "命令笔记 + 实验记录"), (f"Python：{python}", "learn", "1 个 50 行以内脚本"), ("CET-4 与数学竞赛", "英语 45 分钟；数学竞赛 1 小时。11 月后把数学时间转给英语或 CTF。", "英语错题 5 条 + 数学 3 题"), ("整理和输出", "project", "1 次提交或 1 篇短笔记"), ("每日复盘", "review", "复盘卡")],
            2: [(f"网络实验：{net}", "lab", "抓包或拓扑实验记录"), (f"Linux 实验：{linux}", "lab", "排障实验报告"), (f"Web 基础：{web}", "learn", "请求流程图或漏洞原理图"), ("CET-4 与数学竞赛", "英语 45 分钟，数学 1 小时。", "英语错题 5 条"), ("实验报告", "project", "实验报告"), ("每日复盘", "review", "复盘卡")],
            3: [(f"Python：{python}", "learn", "日志提取或自动化脚本"), (f"操作系统与数据库：{security}", "lab", "命令清单和实验记录"), (f"安全基础：{security}", "learn", "5 个概念用自己的话解释"), ("CET-4 与 CTF", "英语 45 分钟；CTF 45 分钟。", "1 道 CTF 记录"), ("CTF Writeup", "project", "300 字 Writeup"), ("每日复盘", "review", "复盘卡")],
            4: [("安全设备与边界", "学习防火墙、IDS/IPS、WAF、VPN、ACL、NAT 的作用和位置。", "设备部署图"), (f"Web 漏洞实验：{web}", "lab", "2 个授权靶场记录"), (f"CTF Web：{ctf}", "learn", "1 道 Web 题记录"), ("CET-4 错题与安全复习", "英语错题重做 45 分钟；复习安全概念 45 分钟。", "错题本更新 10 条"), ("写漏洞报告", "project", "1 份短报告"), ("每日复盘", "review", "复盘卡")],
            5: [(f"项目主线：{python}", "project", "可运行的一步和截图"), (f"项目实验：{linux}", "project", "配置、排错和测试记录"), (f"CTF 专项：{ctf}", "learn", "2 道题记录"), ("CET-4 与 README", "英语 45 分钟，更新项目 README 1 小时。", "文档更新"), ("代码提交和下周待办", "project", "GitHub 提交"), ("每日复盘", "review", "复盘卡")],
            6: [("周测", "网络 30 分钟、Linux 30 分钟、Python 30 分钟、安全 30 分钟。", "周测分数"), (f"项目或 CTF：{blue}", "project", "可展示进展"), ("补最大短板", "只修一个最影响进度的问题。", "问题闭环"), (f"继续项目或 CTF：{ctf}", "project", "1 个输出"), ("周复盘与下周计划", "按 03_复盘系统.md 完成周复盘。", "周复盘表"), ("收尾", "关闭电脑，准备周一休息。", "周一休息确认")],
        }
        return specs.get(wd, [])
    if pidx == 2:
        specs = {
            1: [(f"Web 安全：{web}", "learn", "1 份漏洞原理笔记"), (f"授权靶场：{web}", "lab", "2 个靶场实验记录"), (f"Windows/日志：{security}", "lab", "日志字段和事件记录"), ("CET-4 与 CTF", "英语 45 分钟；CTF 45 分钟。", "英语错题和 1 道 CTF 记录"), ("漏洞报告", "project", "1 份短报告"), ("每日复盘", "review", "复盘卡")],
            2: [(f"日志与流量：{blue}", "lab", "流量或日志分析记录"), (f"Wazuh/ELK：{blue}", "lab", "采集、查询、告警记录"), (f"Python：{python}", "learn", "1 个日志解析脚本"), ("CET-4 与 CTF", "英语 45 分钟；CTF 45 分钟。", "1 道 CTF 记录"), ("事件时间线", "project", "事件时间线"), ("每日复盘", "review", "复盘卡")],
            3: [(f"项目主线：{python}", "project", "项目可运行的一步"), (f"项目实验：{blue}", "project", "测试、截图和结论"), (f"安全基础：{security}", "learn", "5 个概念口述或笔记"), ("CET-4 与 CTF", "英语 45 分钟；CTF 45 分钟。", "1 道 CTF 记录"), ("README 和输出", "project", "文档更新"), ("每日复盘", "review", "复盘卡")],
            4: [(f"Web 综合实验：{web}", "lab", "1 份综合测试报告"), (f"应急响应：{security}", "lab", "1 次告警研判演练"), (f"CTF：{ctf}", "learn", "1 道题和思路复盘"), ("CET-4 与安全复习", "英语 45 分钟，复习漏洞和日志 45 分钟。", "错题本和知识卡片"), ("项目输出", "project", "GitHub 提交"), ("每日复盘", "review", "复盘卡")],
            5: [(f"项目主线：{blue}", "project", "项目中一个可运行模块"), (f"项目实验：{web}", "project", "测试记录和修复建议"), (f"CTF：{ctf}", "learn", "2 道题记录"), ("CET-4 与文档", "英语 45 分钟，写项目文档 1 小时。", "文档更新"), ("代码和提交", "project", "GitHub 提交"), ("每日复盘", "review", "复盘卡")],
            6: [("周测", "Web、日志、Linux、Python 各 30 分钟。", "周测分数"), (f"项目演示：{blue}", "project", "1 个可展示结果"), ("补最大短板", "只修一个最影响项目或找工作的短板。", "问题闭环"), (f"正式 CTF 或模拟赛：{ctf}", "learn", "比赛记录和赛后复盘"), ("周复盘与下周计划", "按 03_复盘系统.md 完成。", "周复盘表"), ("收尾", "准备周一休息。", "周一休息确认")],
        }
        return specs.get(wd, [])
    return []


def advanced_specs(pidx: int, wd: int, python: str, security: str, web: str, blue: str, ctf: str):
    if pidx == 3:
        specs = {
            1: [(f"蓝队专项：{blue}", "lab", "告警研判记录"), (f"安全设备：{security}", "lab", "设备配置或策略检查表"), (f"Python 自动化：{python}", "learn", "1 个自动化工具"), ("英语与面试表达", "英语 45 分钟，整理 3 个面试回答。", "英语记录和 3 个面试回答"), ("简历和项目", "把当天实验写进简历素材库。", "简历素材 1 条"), ("每日复盘", "review", "复盘卡")],
            2: [(f"应急响应：{blue}", "lab", "事件时间线和处置步骤"), (f"安全服务：{security}", "lab", "等保或服务流程笔记"), (f"Web 安全：{web}", "lab", "渗透测试报告片段"), ("英语与面试题", "英语 45 分钟，整理 3 个常见问题。", "面试题库 3 题"), ("项目报告", "把实验转成可展示的项目记录。", "报告更新"), ("每日复盘", "review", "复盘卡")],
            3: [(f"项目主线：{blue}", "project", "项目可运行的一步"), (f"日志与流量：{blue}", "lab", "告警规则或流量分析报告"), (f"CTF：{ctf}", "learn", "1 道题和 Writeup"), ("英语与简历", "英语 45 分钟，更新项目描述和技能标签。", "简历更新记录"), ("模拟面试", "按岗位 JD 进行 30 分钟自问自答。", "面试复盘"), ("每日复盘", "review", "复盘卡")],
            4: [(f"安全项目：{blue}", "project", "项目截图和指标"), (f"漏洞与修复：{web}", "lab", "修复建议和复测记录"), (f"CTF 比赛训练：{ctf}", "learn", "2 道题和赛后复盘"), ("英语与面试题", "英语 45 分钟，准备 3 个项目问题。", "面试回答"), ("求职材料", "更新 GitHub、简历或技术博客。", "材料更新"), ("每日复盘", "review", "复盘卡")],
            5: [(f"项目主线：{blue}", "project", "可展示结果"), (f"项目实验：{web}", "project", "测试和报告更新"), (f"CTF：{ctf}", "learn", "2 道题记录"), ("英语和模拟面试", "英语 45 分钟，模拟面试 1 小时。", "面试复盘"), ("文档和投递准备", "整理岗位关键词、简历版本和项目链接。", "岗位清单"), ("每日复盘", "review", "复盘卡")],
            6: [("周测", "蓝队、Web、Linux、Python、CTF 各 30 分钟。", "周测分数"), (f"项目演示：{blue}", "project", "1 个可展示结果"), ("补最大短板", "只修一个最影响求职的问题。", "问题闭环"), (f"正式比赛或模拟赛：{ctf}", "learn", "比赛记录和赛后复盘"), ("周复盘与下周计划", "按 03_复盘系统.md 完成。", "周复盘表"), ("收尾", "准备周一休息。", "周一休息确认")],
        }
        return specs.get(wd, [])
    if pidx == 4:
        specs = {
            1: [("岗位搜索和 JD 研究", "筛选安全运营、安全服务、等保、应急响应岗位，记录 10 个要求。", "10 个岗位记录"), ("简历定制", "根据岗位关键词修改一版简历，不写虚假经历。", "定制简历"), (f"技能保持：{blue}", "lab", "1 个实验或报告"), ("英语和面试表达", "英语 45 分钟，准备 3 个常见面试问题。", "面试回答 3 条"), ("投递记录", "投递 5～10 个岗位，记录渠道、时间和状态。", "投递表"), ("每日复盘", "review", "复盘卡")],
            2: [("模拟面试", "按真实流程回答 Linux、网络、日志、Web 问题。", "面试复盘"), (f"项目打磨：{blue}", "project", "面试可讲的项目结果"), (f"CTF：{ctf}", "learn", "1 道题和 Writeup"), ("英语和自我介绍", "英语 45 分钟，重写 60 秒自我介绍。", "自我介绍"), ("投递跟进", "跟进回复，记录拒信原因和面试反馈。", "跟进记录"), ("每日复盘", "review", "复盘卡")],
            3: [("面试题库", "整理 20 道安全运营和应急响应问题。", "题库更新"), (f"项目练习：{web}", "lab", "项目演练和讲解录音"), (f"CTF：{ctf}", "learn", "1 道题和赛后总结"), ("英语与岗位研究", "英语 45 分钟，研究 3 家目标公司。", "公司清单"), ("简历和投递", "优化简历、投递并记录。", "投递记录"), ("每日复盘", "review", "复盘卡")],
            4: [(f"蓝队面试题：{blue}", "learn", "口述 10 个问题"), (f"项目报告：{blue}", "project", "报告和演示材料"), (f"CTF：{ctf}", "learn", "1 道题和 Writeup"), ("英语与行为面试", "英语 45 分钟，准备 3 个行为面试故事。", "STAR 回答 3 条"), ("投递和跟进", "投递 5～10 个岗位，更新状态。", "投递表"), ("每日复盘", "review", "复盘卡")],
            5: [("模拟面试复盘", "把本周失败问题重答一遍。", "改进回答"), (f"项目演示：{blue}", "project", "5 分钟讲清一个项目"), (f"CTF：{ctf}", "learn", "2 道题或比赛训练"), ("英语和求职材料", "英语 45 分钟，更新简历、GitHub 和博客。", "材料更新"), ("投递与记录", "投递并记录反馈。", "投递表"), ("每日复盘", "review", "复盘卡")],
            6: [("周测和面试复盘", "复盘岗位要求、面试问题和项目表达。", "周复盘"), (f"项目或 CTF：{ctf}", "project", "可展示结果"), ("补最大短板", "只修一个最影响投递反馈的问题。", "问题闭环"), ("英语和自我介绍", "英语 45 分钟，重录一次自我介绍。", "录音和改稿"), ("周复盘与下周计划", "按 03_复盘系统.md 完成。", "周复盘表"), ("收尾", "准备周一休息。", "周一休息确认")],
        }
        return specs.get(wd, [])
    specs = {
        1: [("实习或工作准备", "检查今天的工作目标，先处理最重要的一件事。", "工作记录"), (f"技能保持：{blue}", "lab", "实验记录"), ("项目沉淀", "把当天工作转成脱敏后的项目描述。", "项目记录"), ("英语和面试", "英语 45 分钟，准备 3 个面试问题。", "面试回答"), ("秋招材料", "更新简历、作品集或投递记录。", "材料更新"), ("每日复盘", "review", "复盘卡")],
        2: [("核心任务", "完成当天最重要的工作或学习任务。", "可验证结果"), (f"安全专项：{security}", "lab", "1 份技术记录"), (f"CTF 或比赛：{ctf}", "learn", "1 道题或比赛记录"), ("英语与面试", "英语 45 分钟，练习项目表达。", "录音或笔记"), ("投递跟进", "更新投递、面试或转正沟通记录。", "跟进记录"), ("每日复盘", "review", "复盘卡")],
        3: [("核心任务", "完成当天最重要的工作或学习任务。", "可验证结果"), (f"项目沉淀：{blue}", "project", "项目文档和证据"), (f"技能训练：{web}", "lab", "1 个实验记录"), ("英语和面试", "英语 45 分钟，准备行为面试故事。", "STAR 回答"), ("文档和复盘", "整理当天工作、问题和下一步。", "工作复盘"), ("每日复盘", "review", "复盘卡")],
        4: [("核心任务", "完成当天最重要的工作或学习任务。", "可验证结果"), (f"技能保持：{blue}", "lab", "1 份实验或报告"), (f"CTF：{ctf}", "learn", "1 道题和 Writeup"), ("英语与面试", "英语 45 分钟，模拟一次面试。", "面试复盘"), ("秋招记录", "更新投递、面试、笔试或转正进度。", "求职表更新"), ("每日复盘", "review", "复盘卡")],
        5: [("核心任务", "完成当天最重要的工作或学习任务。", "可验证结果"), (f"项目沉淀：{blue}", "project", "脱敏项目记录"), (f"CTF：{ctf}", "learn", "1 道题或比赛训练"), ("英语和材料", "英语 45 分钟，更新简历或 GitHub。", "材料更新"), ("周内总结", "记录本周最重要的成果和问题。", "周内总结"), ("每日复盘", "review", "复盘卡")],
        6: [("周测与面试复盘", "复盘岗位要求、面试问题和项目表达。", "周复盘"), (f"项目或 CTF：{ctf}", "project", "可展示结果"), ("补最大短板", "只修一个最影响就业的问题。", "问题闭环"), ("英语和自我介绍", "英语 45 分钟，重录一次自我介绍。", "录音和改稿"), ("周复盘与下周计划", "按 03_复盘系统.md 完成。", "周复盘表"), ("收尾", "准备周一休息。", "周一休息确认")],
    }
    return specs.get(wd, [])


def build_tasks(day: date):
    if day.weekday() == 0:
        return []
    if date(2026, 9, 15) <= day <= date(2026, 9, 20):
        return phase0_tasks(day)
    start, _, _, _ = phase_for(day)
    w = week_index(day, start)
    pidx = phase_index(day)
    wd = day.weekday()
    values = (pidx, wd)
    net = pick(NETWORK_TOPICS, w)
    linux = pick(LINUX_TOPICS, w)
    python = pick(PYTHON_TOPICS, w)
    security = pick(SECURITY_BASE_TOPICS, w)
    web = pick(WEB_TOPICS, w)
    blue = pick(BLUE_TOPICS, w)
    ctf = pick(CTF_TOPICS, w)
    specs = sample_specs(pidx, wd, net, linux, python, security, web, blue, ctf)
    if not specs:
        specs = advanced_specs(pidx, wd, python, security, web, blue, ctf)
    times = ["09:20-11:20", "13:20-15:20", "16:00-17:20", "18:40-20:20", "20:40-21:40", "21:40-22:00"]
    tasks = []
    for idx, spec in enumerate(specs):
        title, mode, output = spec
        if mode in ("learn", "lab", "project", "review"):
            tasks.append(make_row(times[idx], title, mode, output))
        else:
            tasks.append(row(times[idx], title, mode, output))
    return tasks


def three_results(tasks):
    if not tasks:
        return ["完全休息", "不学习、不做 CTF、不安排任务", "周二恢复计划"]
    outputs = []
    for task in tasks:
        if task[3] and task[3] not in outputs:
            outputs.append(task[3])
    while len(outputs) < 3:
        outputs.append("完成当天剩余输出并提交复盘")
    return outputs[:3]


def task_markdown(day: date, tasks):
    start, end, phase_name, focus = phase_for(day)
    wd = WEEKDAYS[day.weekday()]
    lines = [
        f"# {day.isoformat()} {wd} 今日任务", "",
        f"> 阶段：{phase_name}（{start.isoformat()} 至 {end.isoformat()}）",
        f"> 阶段重点：{focus}",
        "> 固定规则：周一完全休息；未经授权的扫描、攻击和访问一律禁止。", "",
        "## 今日三个结果",
    ]
    for idx, item in enumerate(three_results(tasks), 1):
        lines.append(f"{idx}. {item}")
    lines += ["", "## 时间表"]
    if day.weekday() == 0:
        lines += ["", "今天是周一，完全休息。不要学习、不要刷靶场、不要复盘。", "可以做生活、运动、娱乐和休息，周二再恢复。"]
    else:
        lines += ["", "| 时间 | 任务 | 具体怎么做 | 必须产出 |", "|---|---|---|---|"]
        for t, task, how, output in tasks:
            lines.append(f"| {t} | {task} | {how} | {output} |")
        lines += [
            "", "## 开始前 5 分钟", "",
            "1. 打开昨天最后的复盘，读一遍今天的三件事。",
            "2. 关掉短视频、聊天和无关网页，只打开今天第一块所需资料。",
            "3. 先做 10 分钟最小动作，不等状态好再开始。", "",
            "## 卡住时", "",
            "- 10 分钟：写下具体问题，不看答案。",
            "- 30 分钟：查一个可靠来源或问 AI，不能直接抄完整题解。",
            "- 60 分钟：仍无进展，先完成记录，转入下一块，晚上复盘时再处理。", "",
            "## 今日不要做", "",
            "- 不要同时开多个课程和多个方向。",
            "- 不要用看视频代替实验、题目和输出。",
            "- 不要扫描、攻击或访问任何未授权资产。",
            "- 不要因为某一项没完成就熬夜补课。", "",
            "## 最低生存版", "",
            "如果今天有课、生病或突发事务，只完成：CET-4 45 分钟 + 一个核心实验 90 分钟 + 复盘 20 分钟。完成最低版也算达标。",
        ]
    external = load_external_plan(day)
    if external:
        lines += ["", "## 外部 AI 计划补充", "", external, ""]
    return "\n".join(lines) + "\n"


def review_markdown(day: date, tasks):
    start, end, phase_name, focus = phase_for(day)
    wd = WEEKDAYS[day.weekday()]
    lines = [
        f"# {day.isoformat()} {wd} 晚间复盘", "",
        f"> 阶段：{phase_name}",
        "> 21:40 开始，20 分钟内完成。只记录事实，不粉饰。", "",
        "## 今天原计划",
    ]
    if tasks:
        for _, task, _, output in tasks:
            lines.append(f"- {task}，产出：{output}")
    else:
        lines.append("- 周一休息，无学习任务。")
    lines += [
        "", "## 必须回答", "",
        "- 今天真正完成的任务：",
        "- 今天留下的产出物路径或链接：",
        "- 最大卡点：",
        "- 卡点类型：知识不会 / 时间不够 / 情绪问题 / 外界打断",
        "- 未完成原因：",
        "- 今天是否完成 CET-4 45 分钟：是 / 否",
        "- 今天是否完成实验或编程至少 90 分钟：是 / 否",
        "- 明天三件事：",
        "- 明天最重要的一步：",
        "- 今天得分：__ / 10", "",
        "## 评分表", "",
        "| 项目 | 分值 | 自评 | 判断标准 |",
        "|---|---:|---:|---|",
        "| 主任务完成度 | 4 |  | 3 件任务完成得 4 分，完成 2 件得 2 分，低于 2 件得 0 分 |",
        "| 产出物 | 3 |  | 有笔记、代码、报告或 Writeup 得 3 分；只有勾选没有文件得 0 分 |",
        "| 英语 | 1 |  | 完成 45 分钟得 1 分 |",
        "| 健康与休息 | 1 |  | 运动 20 分钟且 23:30 前准备睡觉得 1 分 |",
        "| 诚实执行 | 1 |  | 没有自欺、没有用看视频代替练习得 1 分 |",
        "| **总分** | **10** |  |  |", "",
        "## 复盘后动作", "",
        "- 不低于 8 分：保持结构，不临时加新方向。",
        "- 6～7 分：删掉一个低价值任务，明天先补最重要的短板。",
        "- 低于 6 分：连续 3 天进入最低生存模式，不补双倍任务。",
        "- 连续 2 天未完成：不熬夜，不道歉，不重启计划，只回到今天的第一个核心动作。", "",
        "## 提交格式", "",
        "今日得分：__/10", "今日完成：", "今日产出：", "最大卡点：", "明日三件事：",
    ]
    return "\n".join(lines) + "\n"


def set_run_font(run, ascii_font="Aptos", east_asia="宋体", size=None, color=None, bold=None):
    run.font.name = ascii_font
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def shade_paragraph(paragraph, fill):
    pr = paragraph._p.get_or_add_pPr()
    shd = pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pr.append(shd)
    shd.set(qn("w:fill"), fill)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=50, start=70, bottom=50, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for key, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, "Aptos", "微软雅黑", 8.5, COLORS["muted"])
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(fld)
    run._r.append(instr)
    run._r.append(end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, "Aptos", "微软雅黑", 8.5, COLORS["muted"])


def configure_doc(doc, title):
    doc.core_properties.title = title
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)
    sec.different_first_page_header_footer = True
    add_page_number(sec.footer.paragraphs[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.6)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for name, size, color in [("Title", 24, COLORS["primary"]), ("Heading 1", 16, COLORS["primary"]), ("Heading 2", 13, COLORS["accent"]), ("Heading 3", 11.5, COLORS["primary"])]:
        st = doc.styles[name]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.keep_with_next = True


    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        ppr = doc.styles[style_name]._element.get_or_add_pPr()
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is not None:
            ppr.remove(pbdr)


def add_para(doc, text, quote=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0 if quote else 0.74)
    p.paragraph_format.left_indent = Cm(0.25 if quote else 0)
    p.paragraph_format.right_indent = Cm(0.25 if quote else 0)
    p.paragraph_format.space_after = Pt(5)
    if quote:
        shade_paragraph(p, COLORS["note"])
    r = p.add_run(text)
    set_run_font(r, "Aptos", "宋体", 10.6, COLORS["text"], bold)
    return p


def add_task_table(doc, tasks):
    if not tasks:
        return
    table = doc.add_table(rows=1 + len(tasks), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [2.25, 3.55, 7.1, 3.8]
    for c, text in enumerate(["时间", "任务", "具体怎么做", "必须产出"]):
        cell = table.rows[0].cells[c]
        cell.width = Cm(widths[c])
        shade_cell(cell, COLORS["primary"])
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text)
        set_run_font(r, "Aptos", "微软雅黑", 9.0, "FFFFFF", True)
    for r_idx, task in enumerate(tasks, 1):
        for c_idx, text in enumerate(task):
            cell = table.rows[r_idx].cells[c_idx]
            cell.width = Cm(widths[c_idx])
            set_cell_margins(cell)
            shade_cell(cell, COLORS["alt"] if r_idx % 2 == 0 else "FFFFFF")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_run_font(r, "Aptos", "宋体", 8.9, COLORS["text"], c_idx == 0)
    doc.add_paragraph()


def add_simple_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_data in enumerate(rows):
        for c_idx, text in enumerate(row_data):
            cell = table.rows[r_idx].cells[c_idx]
            cell.width = Cm(widths[c_idx])
            set_cell_margins(cell)
            shade_cell(cell, COLORS["primary"] if r_idx == 0 else (COLORS["alt"] if r_idx % 2 == 0 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, "Aptos", "微软雅黑" if r_idx == 0 else "宋体", 8.9, "FFFFFF" if r_idx == 0 else COLORS["text"], r_idx == 0)
    doc.add_paragraph()


def write_task_docx(day, tasks, path):
    doc = Document()
    configure_doc(doc, f"{day.isoformat()} 今日任务")
    start, end, phase_name, focus = phase_for(day)
    doc.add_heading(f"{day.isoformat()} {WEEKDAYS[day.weekday()]} 今日任务", level=0)
    add_para(doc, f"阶段：{phase_name}（{start.isoformat()} 至 {end.isoformat()}）。阶段重点：{focus}。", quote=True)
    add_para(doc, "固定规则：周一完全休息；未经授权的扫描、攻击和访问一律禁止。", quote=True)
    doc.add_heading("今日三个结果", level=1)
    for idx, result in enumerate(three_results(tasks), 1):
        add_para(doc, f"{idx}. {result}")
    doc.add_heading("时间表", level=1)
    if day.weekday() == 0:
        add_para(doc, "今天是周一，完全休息。不要学习、不要刷靶场、不要复盘。可以做生活、运动、娱乐和休息，周二再恢复。")
    else:
        add_task_table(doc, tasks)
        doc.add_heading("开始前 5 分钟", level=1)
        add_para(doc, "打开昨天最后的复盘，读一遍今天的三件事；关闭无关内容；先做 10 分钟最小动作。")
        doc.add_heading("卡住时", level=1)
        add_para(doc, "10 分钟记录问题，30 分钟查一个可靠来源，60 分钟仍无进展就完成记录并转入下一块。")
        doc.add_heading("今日不要做", level=1)
        add_para(doc, "不要同时开多个课程；不要用看视频代替练习；不要测试未授权资产；不要熬夜补课。")
        doc.add_heading("最低生存版", level=1)
        add_para(doc, "CET-4 45 分钟 + 一个核心实验 90 分钟 + 复盘 20 分钟。完成最低版也算达标。")
    external = load_external_plan(day)
    if external:
        doc.add_heading("外部 AI 计划补充", level=1)
        for part in [p.strip() for p in external.split("\n\n") if p.strip()]:
            add_para(doc, part)
    pointer = path.parent / "current_docx.txt"
    try:
        doc.save(path)
        if pointer.exists():
            pointer.unlink()
    except PermissionError:
        stamp = datetime.now().strftime("%H%M%S")
        fallback = path.with_name(f"{path.stem}_更新版_{stamp}.docx")
        doc.save(fallback)
        pointer.write_text(fallback.name, encoding="utf-8")
    return path


def write_review_docx(day, tasks, path):
    doc = Document()
    configure_doc(doc, f"{day.isoformat()} 晚间复盘")
    start, end, phase_name, focus = phase_for(day)
    doc.add_heading(f"{day.isoformat()} {WEEKDAYS[day.weekday()]} 晚间复盘", level=0)
    add_para(doc, f"阶段：{phase_name}。21:40 开始，20 分钟内完成。只记录事实，不粉饰。", quote=True)
    doc.add_heading("今天原计划", level=1)
    if tasks:
        for _, task, _, output in tasks:
            add_para(doc, f"{task}，产出：{output}")
    else:
        add_para(doc, "周一休息，无学习任务。")
    doc.add_heading("必须回答", level=1)
    questions = [
        "今天真正完成的任务：", "今天留下的产出物路径或链接：", "最大卡点：",
        "卡点类型：知识不会 / 时间不够 / 情绪问题 / 外界打断", "未完成原因：",
        "今天是否完成 CET-4 45 分钟：是 / 否", "今天是否完成实验或编程至少 90 分钟：是 / 否",
        "明天三件事：", "明天最重要的一步：", "今天得分：__ / 10",
    ]
    for q in questions:
        add_para(doc, q)
    doc.add_heading("评分表", level=1)
    add_simple_table(doc, [
        ["项目", "分值", "自评", "判断标准"],
        ["主任务完成度", "4", "", "完成 3 件得 4 分，完成 2 件得 2 分，低于 2 件得 0 分"],
        ["产出物", "3", "", "有笔记、代码、报告或 Writeup 得 3 分"],
        ["英语", "1", "", "完成 45 分钟得 1 分"],
        ["健康与休息", "1", "", "运动 20 分钟且 23:30 前准备睡觉得 1 分"],
        ["诚实执行", "1", "", "没有自欺、没有用看视频代替练习得 1 分"],
        ["总分", "10", "", ""],
    ], [3.2, 1.6, 1.6, 10.0])
    doc.add_heading("复盘后动作", level=1)
    for a in [
        "不低于 8 分：保持结构，不临时加新方向。",
        "6～7 分：删掉一个低价值任务，明天先补最重要的短板。",
        "低于 6 分：连续 3 天进入最低生存模式，不补双倍任务。",
        "连续 2 天未完成：不熬夜，不道歉，不重启计划，只回到今天的第一个核心动作。",
    ]:
        add_para(doc, a)
    doc.add_heading("提交格式", level=1)
    add_para(doc, "今日得分：__/10；今日完成：；今日产出：；最大卡点：；明日三件事：")
    doc.save(path)


def generate(day: date, output_root: Path, mode: str = "both"):
    day_dir = output_root / day.isoformat()
    day_dir.mkdir(parents=True, exist_ok=True)
    tasks = build_tasks(day)
    outputs = []
    if mode in ("morning", "both"):
        md = day_dir / f"{day.isoformat()}_今日任务.md"
        docx = day_dir / f"{day.isoformat()}_今日任务.docx"
        md.write_text(task_markdown(day, tasks), encoding="utf-8")
        write_task_docx(day, tasks, docx)
        outputs.extend([md, docx])
    if mode in ("review", "both"):
        md = day_dir / f"{day.isoformat()}_晚间复盘.md"
        docx = day_dir / f"{day.isoformat()}_晚间复盘.docx"
        md.write_text(review_markdown(day, tasks), encoding="utf-8")
        write_review_docx(day, tasks, docx)
        outputs.extend([md, docx])
    index = output_root / "README.md"
    if not index.exists():
        index.write_text("# 每日任务目录\n\n每天由 generate_daily_task.py 自动生成。\n\n每个日期文件夹包含今日任务和晚间复盘的 Markdown 与 Word 版本。\n", encoding="utf-8")
    return outputs


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", help="YYYY-MM-DD，默认今天")
    parser.add_argument("--mode", choices=["morning", "review", "both"], default="both")
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    args = parser.parse_args()
    day = datetime.strptime(args.date, "%Y-%m-%d").date() if args.date else date.today()
    for path in generate(day, args.output_root, args.mode):
        print(path)


if __name__ == "__main__":
    main()







